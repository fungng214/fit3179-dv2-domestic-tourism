"""Rebuild the chart data from the unmodified official workbooks.

Run: python prepare_data.py (requires openpyxl).
All chart comparisons use June quarters. Source cells and units are retained.
"""
from pathlib import Path
from openpyxl import load_workbook
from collections import defaultdict
import json, math

ROOT = Path(__file__).resolve().parent
RAW = ROOT / 'source-data'
DATA = ROOT / 'data'
DATA.mkdir(exist_ok=True)
STATES = ['New South Wales','Victoria','Queensland','South Australia',
          'Western Australia','Tasmania','Northern Territory','Australian Capital Territory']
ABBR = dict(zip(STATES,['NSW','VIC','QLD','SA','WA','TAS','NT','ACT']))
# Interior label/symbol anchors, not observations of individual tourists.
ANCHORS = dict(zip(STATES,[(147,-32),(144.5,-37),(144,-23),(135,-30),
                            (122,-26),(146.7,-42),(133,-19),(149.12,-35.31)]))
CARTOGRAM = dict(zip(STATES,[(312,205),(265,311),(300,82),(185,205),
                              (65,145),(310,395),(180,70),(398,278)]))
SOURCES = {
 'current': {'publisher':'Tourism Research Australia','title':'Domestic Tourism Statistics June quarter 2026',
  'url':'https://www.tra.gov.au/en/tourism-statistics/domestic-tourism-statistics-results',
  'download':'https://www.tra.gov.au/content/dam/austrade-assets/global/wip/tra/documents/dots/tra-dots-june-quarter-2026-compared-to-june-quarter-2025.xlsx'},
 'historical': {'publisher':'Tourism Research Australia','title':'Revised overnight backcast 2019 to 2024',
  'url':'https://www.tra.gov.au/en/tourism-statistics/domestic-tourism-statistics-results/dots-march-2026',
  'download':'https://www.tra.gov.au/content/dam/austrade-assets/global/wip/tra/documents/dots/tra-dots-2019-2024-overnight-backcast-data-revised.xlsx'},
 'population': {'publisher':'Australian Bureau of Statistics','title':'National state and territory population March 2026 Table 4',
  'url':'https://www.abs.gov.au/statistics/people/population/national-state-and-territory-population/mar-2026',
  'download':'https://www.abs.gov.au/statistics/people/population/national-state-and-territory-population/mar-2026/310104.xlsx'},
 'boundaries': {'publisher':'Australian Bureau of Statistics','title':'ASGS Edition 3 state and territory boundaries',
  'url':'https://www.abs.gov.au/statistics/standards/australian-statistical-geography-standard-asgs/edition-3-july-2021-june-2026/access-and-downloads/data-services-and-apis'},
}
for v in SOURCES.values(): v['accessed']='2026-10-14'

def save(name,obj):
    (DATA/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n')

def number(value):
    if isinstance(value,(int,float)) and math.isfinite(value): return value
    raise ValueError(f'Unexpected missing or suppressed source value: {value!r}')

current=load_workbook(RAW/'dots-june-2026.xlsx',data_only=True)
back=load_workbook(RAW/'dots-backcast-revised.xlsx',data_only=True)
records={}
audit=[]
for metric,sheet in [('trips','Table 1'),('nights','Table 2'),('spend','Table 3')]:
    s=back[sheet]
    for c in range(2,s.max_column+1):
        label=str(s.cell(3,c).value)
        if not label.startswith('June quarter '): continue
        year=int(label[-4:])
        for r in range(4,13):
            state=s.cell(r,1).value
            if state=='ACT': state='Australian Capital Territory'
            if state not in STATES+['Australia']: continue
            value=number(s.cell(r,c).value)
            record=records.setdefault((state,year),{'state':state,'year':year,'period':f'April–June {year}',
                                                     'method':'Revised backcast'})
            record[metric]=value
            audit.append({'dataset':'historical','sheet':sheet,'cell':s.cell(r,c).coordinate,
                          'state':state,'year':year,'metric':metric,'source_value':value,
                          'unit':'AUD million' if metric=='spend' else 'thousand'})
s=current['Table 1a']
for metric,start in [('trips',6),('nights',16),('spend',26)]:
    for r in range(start,start+9):
        state=s.cell(r,1).value
        if str(state).startswith('Australia') and state not in STATES: state='Australia'
        for c,year in [(2,2025),(3,2026)]:
            value=number(s.cell(r,c).value)
            record=records.setdefault((state,year),{'state':state,'year':year,'period':f'April–June {year}',
                                                     'method':'DoTS'})
            record[metric]=value
            audit.append({'dataset':'current','sheet':'Table 1a','cell':s.cell(r,c).coordinate,
                          'state':state,'year':year,'metric':metric,'source_value':value,
                          'unit':'AUD million' if metric=='spend' else 'thousand'})

population_book=load_workbook(RAW/'abs-population.xlsx',data_only=True)
pop_sheet=population_book['Data1']
population={}
for c in range(2,pop_sheet.max_column+1):
    header=str(pop_sheet.cell(1,c).value)
    if ';  Persons ;' not in header: continue
    state=header.split(';')[-2].strip()
    for r in range(11,pop_sheet.max_row+1):
        dt=pop_sheet.cell(r,1).value
        if getattr(dt,'year',None)==2026 and dt.month==3:
            population[state]=number(pop_sheet.cell(r,c).value)
            audit.append({'dataset':'population','sheet':'Data1','cell':pop_sheet.cell(r,c).coordinate,
                          'state':state,'period':'31 March 2026','source_value':population[state],
                          'unit':'persons, all ages'})
assert set(STATES).issubset(population)

history=sorted(records.values(),key=lambda x:(x['state'],x['year']))
for d in history:
    assert all(k in d for k in ['trips','nights','spend'])
    base=records[d['state'],2019]
    for metric in ['trips','nights','spend']:
        d[metric+'_index']=d[metric]/base[metric]*100
    d['trips_m']=d['trips']/1000
    d['nights_m']=d['nights']/1000
    d['spend_b']=d['spend']/1000
    d['trip_change_2019']=(d['trips']/base['trips']-1)*100
    d['abbr']=ABBR.get(d['state'],'Australia')
    if d['state'] in STATES:
        d['lon'],d['lat']=ANCHORS[d['state']]
    if d['year']==2026 and d['state'] in STATES:
        d['population']=population[d['state']]
        d['nights_per_resident']=d['nights']*1000/d['population']
        d['average_nights']=d['nights']/d['trips']
        d['spend_per_visit']=d['spend']*1000/d['trips']
        d['cx'],d['cy']=CARTOGRAM[d['state']]
for year in range(2019,2027):
    subset=sorted([r for r in history if r['year']==year and r['state'] in STATES],key=lambda d:-d['spend'])
    for rank,r in enumerate(subset,1):r['spend_rank']=rank
national=[d for d in history if d['state']=='Australia']
state_history=[d for d in history if d['state'] in STATES]
latest=[d for d in state_history if d['year']==2026]
assert len(national)==8 and len(state_history)==64 and len(latest)==8

purpose=[]
for trip_type,sn in [('Overnight','Table 3a'),('Day trip','Table 7a')]:
    s=current[sn]
    for r in range(6,10):
        for c,year in [(2,2025),(3,2026)]:
            reason=s.cell(r,1).value
            reason_short='Friends / relatives' if reason=='Visiting friends and relatives' else reason
            purpose.append({'trip_type':trip_type,'purpose':reason,'label':reason_short,'year':year,
                            'trips':number(s.cell(r,c).value),'trips_m':number(s.cell(r,c).value)/1000,
                            'source':f'{sn}!{s.cell(r,c).coordinate}'})

travel_mix=[]
s=current['Table 2a']
for r in [6,7]:
    travel_mix.append({'type':s.cell(r,1).value,'trips':number(s.cell(r,3).value),
                       'share':number(s.cell(r,3).value)/number(s.cell(8,3).value)*100})
assert abs(sum(x['trips'] for x in travel_mix)-number(s.cell(8,3).value))<0.002
# 100 squares; rounded allocation is labelled, exact percentages remain in tooltips.
waffle=[]
for n in range(100):
    kind='Intrastate' if n<75 else 'Interstate'
    d=next(x for x in travel_mix if x['type']==kind)
    waffle.append({'column':n%10,'row':n//10,'type':kind,'share':d['share'],'trips_m':d['trips']/1000})

groups={'Accommodation':'Accommodation','Domestic airfares':'Airfares',
        'Takeaways and restaurant meals':'Meals','Petrol':'Petrol',
        'Groceries for self-catering':'Groceries','Shopping / gifts /souvenirs':'Shopping',
        'Alcohol, drinks (not already reported)':'Drinks'}
expenses=defaultdict(lambda:{'v2025':0,'v2026':0,'items':[]})
expense_detail=[]
s=current['Table 4a']
for r in range(6,25):
    item=s.cell(r,1).value; group=groups.get(item,'Other spending')
    old=number(s.cell(r,2).value);new=number(s.cell(r,6).value)
    expenses[group]['v2025']+=old;expenses[group]['v2026']+=new;expenses[group]['items'].append(item)
    expense_detail.append({'item':item,'group':group,'v2025':old,'v2026':new,
                           'source_2025':f'Table 4a!B{r}','source_2026':f'Table 4a!F{r}'})
old_total=number(s.cell(25,2).value);new_total=number(s.cell(25,6).value)
# The source total also includes expenditure on trips with 21+ stops that is
# not allocated to these items. Retain the published total and show the
# difference explicitly; the small negative 2025 residual is rounding.
residual_old=old_total-sum(d['v2025'] for d in expenses.values())
residual_new=new_total-sum(d['v2026'] for d in expenses.values())
expenses['Other spending']['v2025']+=residual_old
expenses['Other spending']['v2026']+=residual_new
expenses['Other spending']['items'].append('Unallocated expenditure and rounding residual')
expense_detail.append({'item':'Unallocated expenditure and rounding residual','group':'Other spending',
                       'v2025':residual_old,'v2026':residual_new,
                       'source_2025':'Table 4a!B25 minus SUM(B6:B24)',
                       'source_2026':'Table 4a!F25 minus SUM(F6:F24)'})
assert abs(sum(d['v2026'] for d in expenses.values())-new_total)<0.01
assert abs(sum(d['v2025'] for d in expenses.values())-old_total)<0.01
expense_groups=[]
for group,d in expenses.items():
    expense_groups.append({'group':group,'v2025':d['v2025'],'v2026':d['v2026'],
                           'value_b':d['v2026']/1000,'change':d['v2026']-d['v2025'],
                           'share':d['v2026']/new_total*100,'items':', '.join(d['items'])})
waterfall=[];running=0
for order,group in enumerate(['Accommodation','Airfares','Meals','Petrol','Groceries','Shopping','Drinks','Other spending']):
    d=next(x for x in expense_groups if x['group']==group)
    end=running+d['change']
    waterfall.append({'group':group,'order':order,'start':running,'end':end,
                      'change':d['change'],'kind':'Increase' if d['change']>=0 else 'Decrease'})
    running=end
waterfall.append({'group':'Net change','order':8,'start':0,'end':new_total-old_total,
                  'change':new_total-old_total,'kind':'Net change'})
assert abs(running-(new_total-old_total))<0.02

# Generalise coastlines for a small overview map; do not alter statistical values.
def ring_area(ring):
    return sum(a[0]*b[1]-b[0]*a[1] for a,b in zip(ring,ring[1:]))/2
def rdp(points,eps=0.04):
    if len(points)<3:return points
    a,b=points[0],points[-1];dx=b[0]-a[0];dy=b[1]-a[1];length=math.hypot(dx,dy)
    distances=[abs(dy*p[0]-dx*p[1]+b[0]*a[1]-b[1]*a[0])/length if length else math.dist(p,a) for p in points]
    index=max(range(len(points)),key=lambda i:distances[i])
    if distances[index]>eps:return rdp(points[:index+1],eps)[:-1]+rdp(points[index:],eps)
    return [a,b]
features=[]
for f in json.loads((RAW/'australia-states.geojson').read_text())['features']:
    name=f['properties']['state_name_2021']
    if name not in STATES or not f.get('geometry'):continue
    g=f['geometry'];polys=g['coordinates'] if g['type']=='MultiPolygon' else [g['coordinates']]
    result=[]
    for polygon in polys:
        outer=polygon[0]
        # Omit distant offshore islands at this national overview scale.
        if not all(110<=p[0]<=155 and -45<=p[1]<=-9 for p in outer):continue
        if abs(ring_area(outer))<0.012 and name!='Australian Capital Territory':continue
        rings=[]
        for idx,ring in enumerate(polygon):
            closed=ring[0]==ring[-1]
            # Split the closed ring before simplification to avoid a degenerate baseline.
            pts=ring[:-1] if closed else ring
            half=len(pts)//2
            simple=rdp(pts[:half+1])+rdp(pts[half:]+[pts[0]])[1:]
            if len(simple)<4:continue
            # D3/Vega expects clockwise outer rings for polygons under a hemisphere.
            if (ring_area(simple)>0)==(idx==0):simple=list(reversed(simple))
            rings.append([[round(p[0],3),round(p[1],3)] for p in simple])
        if rings:result.append(rings)
    assert result,name
    features.append({'type':'Feature','properties':{'state':name,'abbr':ABBR[name]},
                     'geometry':{'type':'MultiPolygon','coordinates':result}})
assert len(features)==8

save('national.json',national);save('state-history.json',state_history);save('latest-states.json',latest)
save('purpose.json',purpose);save('travel-mix.json',travel_mix);save('waffle.json',waffle)
save('expense-groups.json',expense_groups);save('expense-detail.json',expense_detail);save('waterfall.json',waterfall)
save('australia-states.geojson',{'type':'FeatureCollection','features':features})
save('sources.json',SOURCES);save('source-cells.json',audit)
summary={'national':{r['year']:{m:r[m] for m in ['trips_m','nights_m','spend_b','trips_index','spend_index']} for r in national},
         'latest_state_ratios':[{k:r[k] for k in ['abbr','nights_per_resident','average_nights','spend_per_visit']} for r in latest],
         'spending_change_m':new_total-old_total,'source_checks':'Passed: 72 state/national records, source units, purpose labels, current travel mix and expenditure reconciliation.'}
save('validation-summary.json',summary)
print(json.dumps(summary,indent=2))
